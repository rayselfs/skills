#!/usr/bin/env python3
"""
Weekly Report PPTX Generator

Generates or modifies a PowerPoint weekly report from commit data.

Usage:
  # Create from scratch
  python pptx-generator.py --input commits.json --summary summary.json --output report.pptx

  # Modify existing template (uses {{PLACEHOLDER}} markers in slides)
  python pptx-generator.py --template template.pptx --input commits.json --summary summary.json --output report.pptx

Template placeholders (add these text strings inside your .pptx):
  {{SUMMARY}}                    Executive summary paragraph
  {{WEEK}}                       Week number (e.g. 23)
  {{YEAR}}                       Year (e.g. 2025)
  {{STATS}}                      "N commits across N repos"
  {{CATEGORY_Infrastructure}}    Narrative for category named "Infrastructure"
  {{CATEGORY_Features}}          Narrative for category named "Features"
  (any category name works — {{CATEGORY_{Name}}})
"""

import argparse
import json
import sys
from datetime import date, timedelta
from pathlib import Path


def get_week_range():
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    sunday = monday + timedelta(days=6)
    week_num = today.isocalendar()[1]
    year = today.year
    return monday, sunday, week_num, year


def build_replacements(summary_data: dict, stats: dict) -> dict:
    """Build the {{PLACEHOLDER}} → value mapping."""
    monday, sunday, week_num, year = get_week_range()
    replacements = {
        "SUMMARY": summary_data.get("executive_summary", ""),
        "WEEK": str(week_num),
        "YEAR": str(year),
        "STATS": f"{stats['total_commits']} commits across {stats['repos_touched']} repos",
        "DATE_RANGE": f"{monday.strftime('%b %d')} – {sunday.strftime('%b %d, %Y')}",
    }
    for category, narrative in summary_data.get("categories", {}).items():
        key = f"CATEGORY_{category.replace(' ', '_')}"
        replacements[key] = narrative
    return replacements


def replace_in_run(run, replacements: dict):
    """Replace all {{KEY}} occurrences in a single text run."""
    text = run.text
    changed = False
    for key, value in replacements.items():
        marker = f"{{{{{key}}}}}"
        if marker in text:
            text = text.replace(marker, value)
            changed = True
    if changed:
        run.text = text


def apply_replacements_to_slide(slide, replacements: dict):
    """Walk all text runs in a slide and apply replacements."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                replace_in_run(run, replacements)


def modify_template(template_path: str, output_path: str, summary_data: dict, stats: dict):
    """Open an existing .pptx and replace {{PLACEHOLDER}} markers throughout."""
    from pptx import Presentation

    prs = Presentation(template_path)
    replacements = build_replacements(summary_data, stats)

    for slide in prs.slides:
        apply_replacements_to_slide(slide, replacements)

    prs.save(output_path)
    print(f"✓ Template modified → {output_path}")


def create_from_scratch(output_path: str, summary_data: dict, stats: dict):
    """Build a clean weekly report presentation from scratch."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    monday, sunday, week_num, year = get_week_range()
    prs = Presentation()

    # Slide dimensions: widescreen 16:9
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    title_layout = prs.slide_layouts[0]   # Title Slide
    content_layout = prs.slide_layouts[1]  # Title and Content
    blank_layout = prs.slide_layouts[6]    # Blank

    # --- Slide 1: Title ---
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = f"Weekly Report — Week {week_num}, {year}"
    subtitle = slide.placeholders[1]
    subtitle.text = f"{monday.strftime('%B %d')} – {sunday.strftime('%B %d, %Y')}"

    # --- Slide 2: Executive Summary ---
    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Summary"
    tf = slide.placeholders[1].text_frame
    tf.word_wrap = True
    tf.text = summary_data.get("executive_summary", "")

    # --- Slide per category ---
    for category, narrative in summary_data.get("categories", {}).items():
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = category
        tf = slide.placeholders[1].text_frame
        tf.word_wrap = True
        tf.text = narrative

    # --- Last slide: Stats ---
    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Stats"
    tf = slide.placeholders[1].text_frame
    tf.text = ""
    entries = [
        ("Total commits", str(stats["total_commits"])),
        ("Repos touched", str(stats["repos_touched"])),
        ("Week", f"W{week_num} / {year}"),
        ("Period", f"{monday.strftime('%Y-%m-%d')} → {sunday.strftime('%Y-%m-%d')}"),
    ]
    first = True
    for label, value in entries:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = f"{label}: {value}"
        first = False

    prs.save(output_path)
    print(f"✓ Created from scratch → {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Generate or modify a weekly report .pptx from commit data"
    )
    parser.add_argument("--input", required=True, help="Path to commits.json")
    parser.add_argument(
        "--summary",
        required=True,
        help="Path to summary.json (or inline JSON string)",
    )
    parser.add_argument("--output", required=True, help="Output .pptx file path")
    parser.add_argument(
        "--template", help="Optional: existing .pptx template to modify instead of creating from scratch"
    )
    args = parser.parse_args()

    # Load commits
    with open(args.input) as f:
        commits = json.load(f)

    # Load summary — accept file path or raw JSON string
    summary_path = Path(args.summary)
    if summary_path.exists():
        with open(summary_path) as f:
            summary_data = json.load(f)
    else:
        try:
            summary_data = json.loads(args.summary)
        except json.JSONDecodeError:
            print(f"ERROR: --summary must be a valid file path or JSON string", file=sys.stderr)
            sys.exit(1)

    # Compute stats
    repos = set()
    for c in commits:
        parts = [c.get("project", ""), c.get("repo", "")]
        repos.add("/".join(p for p in parts if p))
    stats = {
        "total_commits": len(commits),
        "repos_touched": len(repos),
    }

    if args.template:
        modify_template(args.template, args.output, summary_data, stats)
    else:
        create_from_scratch(args.output, summary_data, stats)


if __name__ == "__main__":
    main()
